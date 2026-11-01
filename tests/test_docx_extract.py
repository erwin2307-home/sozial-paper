from io import BytesIO

from docx import Document

from modules.docx_extract import extract_text


def _build_docx_bytes(paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_text_from_real_docx_succeeds():
    docx_bytes = _build_docx_bytes(["Erster Absatz.", "Zweiter Absatz mit Frage?"])
    text, status = extract_text(docx_bytes)
    assert status == "ok"
    assert "Erster Absatz." in text
    assert "Zweiter Absatz mit Frage?" in text


def test_extract_text_reads_tables_in_document_order():
    document = Document()
    document.add_paragraph("Vor der Tabelle")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Frage in Tabelle?"
    table.cell(0, 1).text = "Antwortfeld"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Zusammengefuehrt?"
    document.add_paragraph("Nach der Tabelle")
    buffer = BytesIO()
    document.save(buffer)
    text, status = extract_text(buffer.getvalue())
    assert status == "ok"
    assert text.splitlines() == [
        "Vor der Tabelle", "Frage in Tabelle? | Antwortfeld", "Zusammengefuehrt?", "Nach der Tabelle",
    ]


def test_extract_text_from_garbage_bytes_fails():
    text, status = extract_text(b"this is not a docx")
    assert status == "failed"
    assert text == ""


def test_extract_text_from_empty_bytes_fails():
    text, status = extract_text(b"")
    assert status == "failed"
    assert text == ""


def test_extract_text_from_empty_docx_fails():
    docx_bytes = _build_docx_bytes([])
    text, status = extract_text(docx_bytes)
    assert status == "failed"
    assert text == ""
