from io import BytesIO
import openpyxl
from modules.export_excel import build_workbook, COLUMNS

def test_build_workbook_has_header_and_rows():
    papers = [
        {"title": "Paper A", "authors": "Max Muster", "year": 2021, "source": "upload",
         "report": {"forschungsfrage": "Frage A", "ergebnisse": "Ergebnis A"}},
        {"title": "Paper B", "authors": "", "year": None, "source": "core", "report": None},
    ]
    xlsx_bytes = build_workbook(papers)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == COLUMNS
    row_a = [cell.value for cell in ws[2]]
    assert row_a[0] == "Paper A"
    assert row_a[COLUMNS.index("forschungsfrage")] == "Frage A"
    row_b = [cell.value for cell in ws[3]]
    assert row_b[0] == "Paper B"
    assert row_b[COLUMNS.index("forschungsfrage")] in ("", None)


def test_build_workbook_sanitizes_formula_injection():
    """Verify formula-like cell values are sanitized to prevent Excel injection attacks."""
    papers = [
        {
            "title": '=HYPERLINK("http://evil.example")',
            "authors": "+IMPORTXML(evil)",
            "year": "-2+5",
            "source": "@SUM(1,1)",
            "report": {
                "forschungsfrage": "=IMPORTXML(A1)",
                "theoretischer_rahmen": "+2+2",
                "methodik": "-1*9",
                "stichprobe": "@DATA",
                "datenerhebung": "normal text",
                "ergebnisse": None,
                "limitationen": "",
                "implikationen": "123 normal",
            },
        }
    ]
    xlsx_bytes = build_workbook(papers)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    row = [cell.value for cell in ws[2]]

    # Verify formula-like prefixes are neutralized by being prefixed with '
    assert row[0].startswith("'="), f"title should be escaped: {repr(row[0])}"
    assert row[1].startswith("'+"), f"authors should be escaped: {repr(row[1])}"
    assert row[2].startswith("'-"), f"year should be escaped: {repr(row[2])}"
    assert row[3].startswith("'@"), f"source should be escaped: {repr(row[3])}"
    assert row[COLUMNS.index("forschungsfrage")].startswith("'="), f"forschungsfrage should be escaped: {repr(row[COLUMNS.index('forschungsfrage')])}"
    assert row[COLUMNS.index("theoretischer_rahmen")].startswith("'+"), f"theoretischer_rahmen should be escaped"
    assert row[COLUMNS.index("methodik")].startswith("'-"), f"methodik should be escaped"
    assert row[COLUMNS.index("stichprobe")].startswith("'@"), f"stichprobe should be escaped"

    # Verify normal text is not modified
    assert row[COLUMNS.index("datenerhebung")] == "normal text", "normal text should not be modified"
    assert row[COLUMNS.index("ergebnisse")] in ("", None), "None should be converted to empty"
    assert row[COLUMNS.index("limitationen")] in ("", None), "empty string should be preserved"
    assert row[COLUMNS.index("implikationen")] == "123 normal", "numeric-starting text should not be modified"


def test_build_workbook_keeps_int_year_as_int_not_stringified():
    papers = [
        {"title": "Paper A", "authors": "Max Muster", "year": 2021, "source": "upload", "report": None},
    ]
    xlsx_bytes = build_workbook(papers)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    row = [cell.value for cell in ws[2]]
    year_value = row[COLUMNS.index("year")]
    assert year_value == 2021
    assert isinstance(year_value, int), f"year should round-trip as int, got {type(year_value)}"


def test_build_workbook_includes_mode_used_column():
    papers = [
        {"title": "Paper A", "authors": "Max Muster", "year": 2021, "source": "upload",
         "mode_used": "B", "report": {"forschungsfrage": "Frage A"}},
    ]
    xlsx_bytes = build_workbook(papers)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    assert "mode_used" in COLUMNS
    header = [cell.value for cell in ws[1]]
    assert header == COLUMNS
    row = [cell.value for cell in ws[2]]
    assert row[COLUMNS.index("mode_used")] == "B"
