from io import BytesIO

from docx import Document

from modules.export_word import build_docx

SECTIONS = [("Abstract", "Absatz eins.\n\nAbsatz zwei."), ("Schlusswort", "Ende ş\x0b mit Steuerzeichen")]


def test_build_docx_contains_title_headings_and_paragraphs_in_order():
    doc = Document(BytesIO(build_docx("Mein Titel", SECTIONS)))
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Mein Titel", "Abstract", "Absatz eins.", "Absatz zwei.", "Schlusswort", "Ende ş mit Steuerzeichen"]
    assert doc.paragraphs[1].style.name == "Heading 1"
