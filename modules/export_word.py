import re
from io import BytesIO

from docx import Document

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: str) -> str:
    return _CONTROL_CHARS.sub("", text or "")


def build_docx(title: str, sections: list) -> bytes:
    document = Document()
    document.add_heading(_clean(title), level=0)
    for heading, body in sections:
        document.add_heading(_clean(heading), level=1)
        for paragraph in re.split(r"\n\s*\n", _clean(body)):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
